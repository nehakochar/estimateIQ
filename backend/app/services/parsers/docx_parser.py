"""
docx_parser.py — Extracts text and tables from Word (.docx) files using python-docx.

A .docx file is a ZIP archive containing XML.  python-docx handles all
the XML parsing and gives us a clean Python API.

Structure of a Word document:
  - Document → Paragraphs (body text, headings, list items)
  - Document → Tables → Rows → Cells → Paragraphs

Our strategy:
  1. Extract all paragraph text from the body.
  2. Extract all tables — converting rows to readable sentences.
     - If the first row looks like a header, use labels: "Header: Value."
     - Otherwise join cells with " | "
  3. Group everything into a single "page 1" entry.
     (Word documents don't have hard page boundaries in the XML.)

Why convert tables to sentences?
  Tab-separated cell dumps are meaningless to the chunking and embedding
  pipeline. Converting "FR-01 | System must support auth | High" into
  "Requirement ID: FR-01. Description: System must support auth. Priority: High."
  gives the classifier and embedder real semantic signal.
"""

import logging
from pathlib import Path

from docx import Document as DocxDocument  # python-docx

from app.services.parsers.base import BaseParser, ParsedPage, ParserError

logger = logging.getLogger(__name__)


def _is_header_row(cells: list[str]) -> bool:
    """
    Return True if a list of cell strings looks like a table header row.

    A row is treated as a header when every non-empty cell:
      - Contains 5 words or fewer (labels, not sentences)
      - Does not end with sentence-terminating punctuation (., ?, !)
    """
    non_empty = [c for c in cells if c]
    if not non_empty:
        return False
    return all(
        len(cell.split()) <= 5 and not cell[-1] in (".", "?", "!")
        for cell in non_empty
    )


class DocxParser(BaseParser):
    """
    Parses .docx files and returns the full document text as a single page.

    Text is extracted from:
      - All body paragraphs (headings, body text, list items, etc.)
      - All tables — converted to readable sentences
    """

    def parse(self, file_path: Path) -> list[ParsedPage]:
        """
        Open the .docx file and extract all text content.

        Args:
            file_path: Path to the .docx file on disk.

        Returns:
            A list containing a single ParsedPage (page=1) with all
            extracted text joined by newlines.

        Raises:
            ParserError: if the file cannot be opened or read.
        """
        logger.info("DocxParser: starting parse of '%s'", file_path)

        try:
            doc = DocxDocument(str(file_path))
        except Exception as exc:
            raise ParserError(
                f"DocxParser: failed to open '{file_path}': {exc}"
            ) from exc

        text_parts: list[str] = []

        try:
            # ── Extract paragraph text ────────────────────────────
            for para in doc.paragraphs:
                stripped = para.text.strip()
                if stripped:
                    text_parts.append(stripped)

            # ── Extract table text ────────────────────────────────
            for table in doc.tables:
                # Visual separator so downstream tools can identify tables
                text_parts.append("\n[TABLE]")

                rows = list(table.rows)
                if not rows:
                    continue

                # Check if first row is a header
                first_row_cells = [cell.text.strip() for cell in rows[0].cells]
                is_header = _is_header_row(first_row_cells)

                if is_header:
                    header_labels = first_row_cells
                    data_rows = rows[1:]
                else:
                    header_labels = []
                    data_rows = rows

                for row in data_rows:
                    cell_texts = [cell.text.strip() for cell in row.cells]

                    # Skip completely empty rows
                    if not any(cell_texts):
                        continue

                    if is_header and header_labels:
                        # Build "Label: Value. Label: Value." sentence
                        parts: list[str] = []
                        for label, value in zip(header_labels, cell_texts):
                            if value:
                                col_name = label if label else "Value"
                                parts.append(f"{col_name}: {value}")
                        if parts:
                            text_parts.append(". ".join(parts) + ".")
                    else:
                        # No header — join non-empty cells with " | "
                        non_empty = [c for c in cell_texts if c]
                        if non_empty:
                            text_parts.append(" | ".join(non_empty))

        except Exception as exc:
            raise ParserError(
                f"DocxParser: error reading content from '{file_path}': {exc}"
            ) from exc

        full_text = "\n".join(text_parts)

        logger.info(
            "DocxParser: finished '%s' — %d characters extracted",
            file_path.name,
            len(full_text),
        )

        return [
            ParsedPage(
                page=1,
                text=full_text,
                metadata={"source": "docx"},
            )
        ]
